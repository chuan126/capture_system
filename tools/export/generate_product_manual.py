#!/usr/bin/env python3
"""从产品使用手册 Markdown 生成 HTML、PDF 和 DOCX 交付文件。"""

from __future__ import annotations

import argparse
import base64
import mimetypes
import re
import sys
import tempfile
from pathlib import Path
from urllib.parse import unquote


def _load_dependencies() -> None:
    try:
        import markdown  # noqa: F401
        import weasyprint  # noqa: F401
        import docx  # noqa: F401
        import cairosvg  # noqa: F401
        import lxml  # noqa: F401
    except ImportError as exc:
        raise SystemExit(
            "缺少文档生成依赖，请安装 markdown、weasyprint、python-docx、cairosvg 和 lxml"
        ) from exc


def _html_template(title: str, body: str) -> str:
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
@page {{
  size: A4;
  margin: 18mm 16mm 18mm 16mm;
  @top-center {{ content: "车载隧道净空高度测量系统产品使用手册"; color: #000; font-size: 8pt; }}
  @bottom-center {{ content: "第 " counter(page) " 页 / 共 " counter(pages) " 页"; color: #000; font-size: 8pt; }}
}}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #000; font-family: "WenQuanYi Micro Hei", "Microsoft YaHei", sans-serif; font-size: 10.2pt; line-height: 1.72; }}
h1 {{ margin: 54mm 0 12mm; color: #000; font-size: 25pt; line-height: 1.35; text-align: center; }}
h2 {{ margin: 10mm 0 4mm; color: #000; border-bottom: 1.5px solid #777; padding-bottom: 2mm; font-size: 17pt; line-height: 1.35; break-after: avoid; }}
h2.chapter {{ break-before: page; }}
h3 {{ margin: 7mm 0 2.5mm; color: #000; font-size: 13pt; line-height: 1.4; break-after: avoid; }}
h4 {{ margin: 5mm 0 2mm; color: #000; font-size: 11pt; break-after: avoid; }}
p {{ margin: 0 0 3mm; text-align: justify; }}
body > p:nth-of-type(-n+4) {{ text-align: center; }}
hr {{ border: 0; border-top: 1px solid #aaa; margin: 8mm 0; }}
hr.cover-break, hr.toc-break {{ visibility: hidden; height: 0; margin: 0; break-after: page; }}
a {{ color: #000; text-decoration: underline; }}
code {{ padding: 0.2mm 1mm; border-radius: 2px; background: #f1f1f1; color: #000; font-family: "DejaVu Sans Mono", monospace; font-size: 9pt; }}
pre {{ margin: 3mm 0 4mm; padding: 4mm; border: 1px solid #aaa; border-radius: 4px; background: #f7f7f7; white-space: pre-wrap; break-inside: avoid; }}
pre code {{ padding: 0; background: transparent; }}
blockquote {{ margin: 4mm 0; padding: 3.5mm 4mm; border-left: 4px solid #666; background: #f5f5f5; color: #000; break-inside: avoid; }}
blockquote p {{ margin: 0 0 2mm; }}
blockquote p:last-child {{ margin-bottom: 0; }}
ul, ol {{ margin: 1.5mm 0 3mm; padding-left: 7mm; }}
li {{ margin-bottom: 1.2mm; }}
table {{ width: 100%; margin: 3mm 0 5mm; border-collapse: collapse; font-size: 8.6pt; break-inside: auto; }}
thead {{ display: table-header-group; }}
tr {{ break-inside: avoid; }}
th {{ background: #ededed; color: #000; font-weight: 700; }}
th, td {{ border: 1px solid #999; padding: 2mm 2.2mm; vertical-align: top; text-align: left; }}
img {{ display: block; max-width: 100%; max-height: 205mm; margin: 4mm auto 2mm; object-fit: contain; break-inside: avoid; }}
p > em:only-child {{ display: block; color: #000; font-size: 8.6pt; text-align: center; }}
.toc {{ margin: 4mm 0 6mm; padding: 4mm 7mm; border: 1px solid #aaa; border-radius: 4px; background: #fafafa; }}
.toc ul {{ list-style: none; padding-left: 0; }}
.toc ul ul {{ padding-left: 6mm; }}
.toc li {{ margin: 1mm 0; }}
.toc > ul > li > a {{ font-weight: 700; }}
@media screen {{ body {{ max-width: 980px; margin: 0 auto; padding: 30px; background: white; }} }}
</style>
</head>
<body>
{body}
</body>
</html>
"""


def _mark_pdf_breaks(body_html: str) -> str:
    from lxml import html

    root = html.fragment_fromstring(body_html, create_parent="div")
    horizontal_rules = root.xpath("./hr")
    if horizontal_rules:
        horizontal_rules[0].set("class", "cover-break")
    if len(horizontal_rules) > 1:
        horizontal_rules[1].set("class", "toc-break")
    for heading in root.xpath(".//h2"):
        text = "".join(heading.itertext()).strip()
        if re.match(r"^(?:\d+\s|附录)", text):
            heading.set("class", "chapter")
    return "".join(html.tostring(child, encoding="unicode") for child in root)


def _set_run_font(run, name: str = "Microsoft YaHei", size_pt: float | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _add_inline(paragraph, node) -> None:
    from lxml import etree

    if isinstance(node, str):
        if node:
            _set_run_font(paragraph.add_run(node))
        return
    if not isinstance(node.tag, str):
        return
    tag = node.tag.lower()
    if tag == "br":
        paragraph.add_run().add_break()
    elif tag == "a":
        label = "".join(node.itertext()).strip()
        href = node.get("href", "")
        display = label if href.startswith("#") or not href else f"{label}（{href}）"
        run = paragraph.add_run(display)
        _set_run_font(run)
        run.font.underline = True
    else:
        before = len(paragraph.runs)
        if node.text:
            _add_inline(paragraph, node.text)
        for child in node:
            _add_inline(paragraph, child)
            if child.tail:
                _add_inline(paragraph, child.tail)
        for run in paragraph.runs[before:]:
            if tag in {"strong", "b"}:
                run.bold = True
            if tag in {"em", "i"}:
                run.italic = True
            if tag == "code":
                _set_run_font(run, "Consolas", 9)


def _add_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph.alignment = 1
    prefix = paragraph.add_run("第 ")
    _set_run_font(prefix, size_pt=8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for field_node in (begin, instr, separate):
        field_run = paragraph.add_run()
        _set_run_font(field_run, size_pt=8)
        field_run._r.append(field_node)
    result_run = paragraph.add_run("1")
    _set_run_font(result_run, size_pt=8)
    end_run = paragraph.add_run()
    _set_run_font(end_run, size_pt=8)
    end_run._r.append(end)
    suffix = paragraph.add_run(" 页")
    _set_run_font(suffix, size_pt=8)


def _render_svg_for_docx(svg_path: Path, output_path: Path) -> None:
    """将SVG及其本地引用图片合成为单张PNG，避免Word丢失外部底图。"""
    import cairosvg
    from lxml import etree

    root = etree.parse(str(svg_path)).getroot()
    xlink_href = "{http://www.w3.org/1999/xlink}href"
    for image_element in root.xpath("//*[local-name()='image']"):
        href_attribute = "href" if image_element.get("href") else xlink_href
        href = image_element.get(href_attribute)
        if not href or href.startswith(("data:", "http://", "https://")):
            continue
        referenced_path = (svg_path.parent / unquote(href)).resolve()
        if not referenced_path.is_file():
            raise FileNotFoundError(f"SVG引用的底图不存在：{referenced_path}")
        mime_type = mimetypes.guess_type(referenced_path.name)[0] or "application/octet-stream"
        encoded = base64.b64encode(referenced_path.read_bytes()).decode("ascii")
        image_element.set(href_attribute, f"data:{mime_type};base64,{encoded}")

    cairosvg.svg2png(
        bytestring=etree.tostring(root),
        write_to=str(output_path),
        output_width=1920,
    )


def _build_docx(body_html: str, source_dir: Path, output_path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
    from lxml import html
    from PIL import Image

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("车载隧道净空高度测量系统产品使用手册")
    _set_run_font(header_run, size_pt=8)
    _add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    for style_name in ["Normal", "Body Text", "List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
    for level in range(1, 5):
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)

    root = html.fragment_fromstring(body_html, create_parent="div")
    image_cache: dict[Path, Path] = {}

    def add_image(src: str, alt: str) -> None:
        image_path = (source_dir / unquote(src)).resolve()
        if not image_path.exists():
            paragraph = document.add_paragraph(f"[图片缺失：{alt or src}]")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        ready_path = image_path
        if image_path.suffix.lower() == ".svg":
            cached_path = image_cache.get(image_path)
            if cached_path is None:
                ready_path = Path(tempfile.gettempdir()) / f"capture-manual-{len(image_cache)}.png"
                _render_svg_for_docx(image_path, ready_path)
                image_cache[image_path] = ready_path
            else:
                ready_path = cached_path
        try:
            with Image.open(ready_path) as opened:
                pixel_width = opened.width
            width = min(6.5, max(4.3, pixel_width / 250.0))
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(ready_path), width=Inches(width))
        except Exception as exc:  # pragma: no cover - 仅在交付机图片损坏时触发
            document.add_paragraph(f"[图片无法插入：{alt or src}，{exc}]")

    def add_table(element) -> None:
        rows = element.xpath(".//tr")
        if not rows:
            return
        column_count = max(len(row.xpath("./th|./td")) for row in rows)
        table = document.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            cells = row.xpath("./th|./td")
            for column_index, source_cell in enumerate(cells):
                target = table.cell(row_index, column_index)
                target.text = ""
                paragraph = target.paragraphs[0]
                _add_inline(paragraph, source_cell)
                for run in paragraph.runs:
                    _set_run_font(run, size_pt=8.5)
                    run.bold = row_index == 0 and source_cell.tag.lower() == "th"
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "EDEDED")
                    target._tc.get_or_add_tcPr().append(shading)

    def process(element, list_style: str | None = None) -> None:
        if not isinstance(element.tag, str):
            return
        tag = element.tag.lower()
        if tag in {"h1", "h2", "h3", "h4"}:
            level = int(tag[1])
            paragraph = document.add_heading(level=level)
            _add_inline(paragraph, element)
            text = "".join(element.itertext()).strip()
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(180)
            if level == 2 and (re.match(r"^\d+\s", text) or text.startswith("附录")):
                paragraph.paragraph_format.page_break_before = True
            return
        if tag == "p":
            images = element.xpath("./img")
            if len(images) == 1 and not "".join(element.itertext()).strip():
                add_image(images[0].get("src", ""), images[0].get("alt", ""))
                return
            paragraph = document.add_paragraph(style=list_style)
            _add_inline(paragraph, element)
            paragraph.paragraph_format.space_after = Pt(5)
            return
        if tag == "blockquote":
            for child in element:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.right_indent = Cm(0.2)
                _add_inline(paragraph, child)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
            return
        if tag in {"ul", "ol"}:
            style = "List Bullet" if tag == "ul" else "List Number"
            for item in element.xpath("./li"):
                paragraph = document.add_paragraph(style=style)
                _add_inline(paragraph, item)
            return
        if tag == "table":
            add_table(element)
            return
        if tag == "pre":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            run = paragraph.add_run("".join(element.itertext()))
            _set_run_font(run, "Consolas", 8.5)
            return
        if tag == "hr":
            return
        for child in element:
            process(child)

    for child in root:
        process(child)

    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="产品使用手册 Markdown 文件")
    parser.add_argument("--output-dir", type=Path, help="输出目录，默认与源文件相同")
    args = parser.parse_args()
    _load_dependencies()

    import markdown
    from weasyprint import HTML

    source = args.source.resolve()
    if not source.is_file():
        raise SystemExit(f"源文件不存在：{source}")
    output_dir = (args.output_dir or source.parent).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    markdown_text = source.read_text(encoding="utf-8")
    renderer = markdown.Markdown(
        extensions=["tables", "fenced_code", "toc"],
        extension_configs={"toc": {"permalink": False, "toc_depth": "2-3"}},
    )
    body = _mark_pdf_breaks(renderer.convert(markdown_text))
    html_text = _html_template("车载隧道净空高度测量系统产品使用手册", body)

    html_path = output_dir / "产品使用手册.html"
    pdf_path = output_dir / "产品使用手册.pdf"
    docx_path = output_dir / "产品使用手册.docx"
    html_path.write_text(html_text, encoding="utf-8")
    HTML(filename=str(html_path), base_url=str(source.parent)).write_pdf(str(pdf_path))
    _build_docx(body, source.parent, docx_path)

    print(f"HTML: {html_path}")
    print(f"PDF: {pdf_path}")
    print(f"DOCX: {docx_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
